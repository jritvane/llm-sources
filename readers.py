"""Lähteiden lukeminen: jokainen tiedostomuoto muunnetaan tekstiksi."""
import csv
import io
import sys
from pathlib import Path
from urllib.parse import urlparse

import pymupdf
import requests
from bs4 import BeautifulSoup
from docx import Document

HTTP_TIMEOUT_SECONDS = 20
HTTP_HEADERS = {"User-Agent": "llm-sources/1.0 (course assignment)"}
TEXT_EXTENSIONS = {".txt", ".md"}
HTML_EXTENSIONS = {".html", ".htm"}


class SourceError(Exception):
    """Lähdettä ei voitu lukea. Viesti on tarkoitettu käyttäjälle."""


def is_url(source: str) -> bool:
    return urlparse(source).scheme in ("http", "https")


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceError(f"{path}: tiedoston merkistöä ei tunnistettu.")


def html_to_text(html: str | bytes) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript", "template", "svg",
                         "nav", "footer", "aside", "form"]):
        element.decompose()
    content = soup.find("main") or soup.find("article") or soup.body or soup
    return content.get_text("\n", strip=True)


def read_csv_file(path: Path) -> str:
    text = read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text, newline=""), dialect)
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if row)


def read_docx_file(path: Path) -> str:
    try:
        document = Document(str(path))
    except Exception as e:  # Kirjasto heittää useita eri poikkeustyyppejä
        raise SourceError(f"{path}: Word-tiedostoa ei voitu avata ({e}).") from e

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, start=1):
        parts.append(f"[Taulukko {number}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def read_pdf_file(path: Path) -> str:
    try:
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise SourceError(f"{path}: PDF on salasanasuojattu.")
            page_texts = [page.get_text() for page in document]
    except SourceError:
        raise
    except Exception as e:  # Vioittunut tai muu kuin PDF-tiedosto
        raise SourceError(f"{path}: PDF-tiedostoa ei voitu avata ({e}).") from e

    if not any(text.strip() for text in page_texts):
        raise SourceError(f"{path}: PDF:ssä ei ole tekstiä. Se on ehkä skannattu kuva.")
    return "\n".join(f"[Sivu {n}]\n{text}" for n, text in enumerate(page_texts, start=1))


def read_url(url: str) -> str:
    try:
        response = requests.get(url, timeout=HTTP_TIMEOUT_SECONDS, headers=HTTP_HEADERS)
        response.raise_for_status()
    except requests.exceptions.Timeout as e:
        raise SourceError(f"{url}: palvelin ei vastannut {HTTP_TIMEOUT_SECONDS} sekunnissa.") from e
    except requests.exceptions.ConnectionError as e:
        raise SourceError(f"{url}: palvelimeen ei saatu yhteyttä. Tarkista osoite ja verkkoyhteys.") from e
    except requests.exceptions.HTTPError as e:
        raise SourceError(f"{url}: palvelin palautti virheen {e.response.status_code}.") from e
    except requests.exceptions.RequestException as e:
        raise SourceError(f"{url}: sivua ei voitu hakea ({e}).") from e

    content_type = response.headers.get("Content-Type", "")
    if "html" in content_type:
        return html_to_text(response.content)
    if content_type.startswith("text/"):
        return response.text
    raise SourceError(f"{url}: sisältötyyppiä '{content_type}' ei tueta.")


def read_source(source: str) -> str:
    if is_url(source):
        return read_url(source)

    path = Path(source)
    if not path.is_file():
        raise SourceError(f"{source}: tiedostoa ei löydy.")

    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return read_text_file(path)
    if suffix in HTML_EXTENSIONS:
        return html_to_text(path.read_bytes())
    if suffix == ".csv":
        return read_csv_file(path)
    if suffix == ".docx":
        return read_docx_file(path)
    if suffix == ".pdf":
        return read_pdf_file(path)
    raise SourceError(f"{source}: tiedostomuotoa '{suffix or '(ei päätettä)'}' ei tueta.")


if __name__ == "__main__":
    # Testikäyttö: python readers.py tiedosto1 tiedosto2 ...
    for source in sys.argv[1:]:
        try:
            text = read_source(source)
        except SourceError as e:
            print(f"VIRHE: {e}\n")
            continue
        print(f"=== {source}: {len(text)} merkkiä ===")
        print(text[:300])
        print("...\n")